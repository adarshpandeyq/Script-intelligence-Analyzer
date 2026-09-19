"""Generate the project presentation for Script Intelligence Analyzer.

Run:  python generate_ppt.py
Out:  ../public/downloads/Script_Intelligence_Analyzer.pptx
"""
from __future__ import annotations

import os
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

BG = RGBColor(0x0A, 0x09, 0x18)
PANEL = RGBColor(0x16, 0x14, 0x33)
ACCENT = RGBColor(0x8B, 0x5C, 0xF6)
ACCENT2 = RGBColor(0x38, 0xBD, 0xF8)
TEXT = RGBColor(0xE9, 0xE7, 0xFF)
MUTED = RGBColor(0xA5, 0xA0, 0xC8)

OUT_DIR = os.path.join(os.path.dirname(__file__), "..", "public", "downloads")
OUT_FILE = os.path.join(OUT_DIR, "Script_Intelligence_Analyzer.pptx")


def add_bg(slide, prs):
    bg = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    bg.shadow.inherit = False
    glow = slide.shapes.add_shape(1, Inches(-2), Inches(-2.5), Inches(6), Inches(6))
    glow.fill.solid()
    glow.fill.fore_color.rgb = RGBColor(0x2A, 0x1B, 0x5E)
    glow.line.fill.background()
    glow.shadow.inherit = False
    glow2 = slide.shapes.add_shape(1, Inches(8.6), Inches(3.6), Inches(6), Inches(6))
    glow2.fill.solid()
    glow2.fill.fore_color.rgb = RGBColor(0x11, 0x2C, 0x54)
    glow2.line.fill.background()
    glow2.shadow.inherit = False


def add_title(slide, title, kicker=None):
    if kicker:
        box = slide.shapes.add_textbox(Inches(0.7), Inches(0.42), Inches(11.9), Inches(0.4))
        p = box.text_frame.paragraphs[0]
        p.text = kicker.upper()
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = ACCENT2
    box = slide.shapes.add_textbox(Inches(0.7), Inches(0.82 if kicker else 0.55), Inches(11.9), Inches(0.9))
    p = box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = TEXT
    bar = slide.shapes.add_shape(1, Inches(0.72), Inches(1.78 if kicker else 1.5), Inches(1.6), Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    bar.shadow.inherit = False


def add_bullets(slide, items, top=2.15, left=0.8, width=11.4, size=16):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(5.0))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "▸  " + item
        p.font.size = Pt(size)
        p.font.color.rgb = TEXT
        p.space_after = Pt(12)
    return box


def add_card(slide, left, top, width, height, heading, body, accent=ACCENT):
    shape = slide.shapes.add_shape(5, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = PANEL
    shape.line.color.rgb = accent
    shape.line.width = Pt(1)
    shape.shadow.inherit = False
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.22)
    tf.margin_top = Inches(0.16)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.text = heading
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = accent
    p2 = tf.add_paragraph()
    p2.text = body
    p2.font.size = Pt(11.5)
    p2.font.color.rgb = MUTED
    p2.space_before = Pt(6)
    return shape


def add_footer(slide, index):
    box = slide.shapes.add_textbox(Inches(0.7), Inches(6.75), Inches(11.9), Inches(0.35))
    p = box.text_frame.paragraphs[0]
    p.text = f"Script Intelligence Analyzer  •  AI / NLP Project  •  {index:02d}"
    p.font.size = Pt(10)
    p.font.color.rgb = MUTED
    p.alignment = PP_ALIGN.RIGHT


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1 — Title
    s = prs.slides.add_slide(blank)
    add_bg(s, prs)
    chip = s.shapes.add_shape(5, Inches(0.9), Inches(1.9), Inches(3.5), Inches(0.55))
    chip.fill.solid(); chip.fill.fore_color.rgb = PANEL
    chip.line.color.rgb = ACCENT; chip.shadow.inherit = False
    cp = chip.text_frame.paragraphs[0]
    cp.text = "AI · NLP · CINEMATIC ANALYTICS"; cp.font.size = Pt(12)
    cp.font.color.rgb = ACCENT2; cp.alignment = PP_ALIGN.CENTER
    t = s.shapes.add_textbox(Inches(0.9), Inches(2.6), Inches(11.5), Inches(1.4))
    tp = t.text_frame.paragraphs[0]
    tp.text = "Script Intelligence Analyzer"
    tp.font.size = Pt(52); tp.font.bold = True; tp.font.color.rgb = TEXT
    sub = s.shapes.add_textbox(Inches(0.9), Inches(4.05), Inches(11.5), Inches(0.8))
    sp = sub.text_frame.paragraphs[0]
    sp.text = "Upload a movie or short-film script. Get characters, scenes, emotions, themes,\nsuspense curves and foreshadowing — automatically."
    sp.font.size = Pt(18); sp.font.color.rgb = MUTED
    meta = s.shapes.add_textbox(Inches(0.9), Inches(5.5), Inches(11.5), Inches(0.8))
    mp = meta.text_frame.paragraphs[0]
    mp.text = "Python · Flask · spaCy · NLTK · Transformers · Sentence-Transformers · scikit-learn · PyMuPDF · NetworkX · Plotly"
    mp.font.size = Pt(13); mp.font.color.rgb = ACCENT2
    add_footer(s, 1)

    # 2 — Problem statement
    s = prs.slides.add_slide(blank)
    add_bg(s, prs); add_title(s, "Problem Statement", "01 — Why this project")
    add_bullets(s, [
        "Screenplays are unstructured: scene headings, character cues, parentheticals and dialogue are hard to parse at scale.",
        "Manual script coverage is slow and subjective — studios and students spend hours on one draft.",
        "There is no quick way to quantify tension over time or to see whether early scenes set up later payoffs.",
        "Emotion, theme and character prominence are inferred by reading, not measured.",
        "Writers need fast, objective feedback on pacing, suspense peaks and theme consistency.",
    ])

    # 3 — Objectives
    s = prs.slides.add_slide(blank)
    add_bg(s, prs); add_title(s, "Project Objectives", "02 — Goals")
    cards = [
        ("Automate parsing", "Robust extraction from PDF & TXT with industry screenplay formatting rules."),
        ("Detect characters", "Identify cast, dialogue volume and the protagonist automatically."),
        ("Segment scenes", "Split into scenes with location, time, cast and key action."),
        ("Quantify emotion", "Score fear, joy, sadness, anger, surprise and neutral per scene."),
        ("Track suspense", "Build a scene-by-scene tension curve and flag high-tension peaks."),
        ("Find foreshadowing", "Link early setups to later payoffs via semantic similarity."),
    ]
    for i, (h, b) in enumerate(cards):
        col, row = i % 3, i // 3
        add_card(s, 0.8 + col * 4.05, 2.3 + row * 2.2, 3.7, 1.85, h, b,
                 ACCENT if i % 2 == 0 else ACCENT2)
    add_footer(s, 3)

    # 4 — Proposed solution
    s = prs.slides.add_slide(blank)
    add_bg(s, prs); add_title(s, "Proposed Solution", "03 — Approach")
    add_bullets(s, [
        "A modular Python NLP pipeline: upload → extract → clean → analyse → report.",
        "Rule-based screenplay parser (headings, cues, transitions, parentheticals) tuned for real scripts.",
        "Pretrained models instead of training from scratch: transformer embeddings for semantic similarity, zero-shot classification for themes/emotions.",
        "Lexicon + model hybrid scoring keeps results fast, explainable and offline-friendly.",
        "Interactive web UI (Flask + Plotly) with glassmorphism cards and cinematic dark theme.",
        "Every result is stored so analyses can be revisited and compared.",
    ])
    add_footer(s, 4)

    # 5 — Architecture
    s = prs.slides.add_slide(blank)
    add_bg(s, prs); add_title(s, "System Architecture", "04 — Pipeline")
    flow = ["Upload\nPDF / TXT", "Text\nExtraction", "Cleaning &\nNormalisation", "NLP\nModules", "Intelligence\nReport"]
    for i, label in enumerate(flow):
        box = s.shapes.add_shape(5, Inches(0.6 + i * 2.55), Inches(2.7), Inches(2.0), Inches(1.35))
        box.fill.solid(); box.fill.fore_color.rgb = PANEL
        box.line.color.rgb = ACCENT if i % 2 == 0 else ACCENT2
        box.shadow.inherit = False
        tf = box.text_frame; tf.word_wrap = True
        tf.paragraphs[0].text = label
        tf.paragraphs[0].font.size = Pt(13)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = TEXT
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        if i < len(flow) - 1:
            ar = s.shapes.add_textbox(Inches(2.62 + i * 2.55), Inches(3.15), Inches(0.5), Inches(0.5))
            ap = ar.text_frame.paragraphs[0]
            ap.text = "➜"; ap.font.size = Pt(20); ap.font.color.rgb = ACCENT2
    add_card(s, 0.6, 4.5, 3.9, 1.6, "NLP Modules", "Characters · Scenes · Emotion · Themes · Suspense · Foreshadowing", ACCENT)
    add_card(s, 4.75, 4.5, 3.9, 1.6, "Storage", "Analysis records + JSON payloads (PostgreSQL / SQLite)", ACCENT2)
    add_card(s, 8.9, 4.5, 3.9, 1.6, "Presentation", "Flask + Plotly graphs, glassmorphism UI, JSON API", ACCENT)
    add_footer(s, 5)

    # 6 — Technologies
    s = prs.slides.add_slide(blank)
    add_bg(s, prs); add_title(s, "NLP Technologies Used", "05 — Stack")
    tech = [
        ("spaCy", "Tokenisation, POS tagging, named entities, sentence segmentation."),
        ("NLTK", "Stop-words, tokenisation and lexicon utilities."),
        ("Transformers", "Zero-shot classification for emotion & theme labelling."),
        ("Sentence-Transformers", "Dense embeddings for scene similarity / foreshadowing."),
        ("scikit-learn", "TF-IDF vectors, cosine similarity, clustering."),
        ("PyMuPDF", "Fast and reliable PDF text extraction."),
        ("NetworkX", "Character co-occurrence / interaction graphs."),
        ("Plotly + Flask", "Interactive charts and the web application layer."),
    ]
    for i, (h, b) in enumerate(tech):
        col, row = i % 2, i // 2
        add_card(s, 0.8 + col * 6.1, 2.3 + row * 1.28, 5.7, 1.1, h, b,
                 ACCENT if i % 2 == 0 else ACCENT2)
    add_footer(s, 6)

    # 7 — Features
    s = prs.slides.add_slide(blank)
    add_bg(s, prs); add_title(s, "Main Features", "06 — What it delivers")
    add_bullets(s, [
        "Script statistics: words, dialogue lines, scenes, characters, estimated runtime.",
        "Character intelligence: dialogue counts, presence map, protagonist detection.",
        "Scene analysis: location, time of day, cast and an AI-generated mini summary.",
        "Emotion analysis with intensity per scene and an interactive stacked/scroll chart.",
        "Theme detection with relevance scores and supporting keywords.",
        "Suspense & tension meter with peak-scene highlighting for horror/thriller scripts.",
        "Possible foreshadowing pairs with explanations and confidence levels.",
        "One-click final intelligence report + downloadable JSON export.",
    ], size=15)
    add_footer(s, 7)

    # 8 — Workflow
    s = prs.slides.add_slide(blank)
    add_bg(s, prs); add_title(s, "Analysis Workflow", "07 — Step by step")
    steps = [
        ("1", "Upload", "User submits a .txt or .pdf screenplay (with validation & progress)."),
        ("2", "Extract", "PyMuPDF / plain-text reader pulls raw text, then cleaning removes page numbers and artefacts."),
        ("3", "Parse", "Screenplay parser separates scene headings, action, character cues, parentheticals and dialogue."),
        ("4", "Analyse", "Six independent NLP modules run over the parsed structure."),
        ("5", "Visualise", "Plotly charts render the emotion mix, suspense curve and character presence."),
        ("6", "Report", "The report generator merges everything into a concise intelligence brief."),
    ]
    for i, (n, h, b) in enumerate(steps):
        col, row = i % 2, i // 2
        add_card(s, 0.8 + col * 6.1, 2.25 + row * 1.5, 5.7, 1.25, f"Step {n} · {h}", b,
                 ACCENT if i % 2 == 0 else ACCENT2)
    add_footer(s, 8)

    # 9 — Expected results
    s = prs.slides.add_slide(blank)
    add_bg(s, prs); add_title(s, "Expected Results", "08 — Demo outcome")
    add_bullets(s, [
        "Upload a 10-scene horror short → the pipeline returns 8–12 detected characters and 10 segmented scenes in seconds.",
        "The protagonist is correctly ranked by dialogue volume, presence and scene coverage.",
        "Suspense curve rises toward the climax; peak scenes are flagged with a tension score > 70.",
        "Emotion profile of a horror script is dominated by fear, surprise and neutral exposition.",
        "Foreshadowing module links the cold-open setup to the final revelation (labelled POSSIBLE).",
        "Final report summarises stats, cast, emotions, themes, tension arc and possible foreshadowing.",
    ], size=15)
    add_footer(s, 9)

    # 10 — Future scope
    s = prs.slides.add_slide(blank)
    add_bg(s, prs); add_title(s, "Future Scope", "09 — Roadmap")
    add_bullets(s, [
        "Beat / act structure detection (Save the Cat, three-act) with pacing suggestions.",
        "Character relationship graphs with sentiment-weighted interactions (NetworkX).",
        "Fine-tuned transformer models for screenplay-specific emotion & theme labels.",
        "Dialogue quality scoring, rewrite suggestions and automatic loglines.",
        "Scene-to-shot breakdown with estimated screen time and budget hints.",
        "Comparison mode: track how tension and themes change between draft versions.",
    ], size=15)
    add_footer(s, 10)

    os.makedirs(OUT_DIR, exist_ok=True)
    prs.save(OUT_FILE)
    print(f"Saved presentation -> {os.path.abspath(OUT_FILE)}")


if __name__ == "__main__":
    build()
